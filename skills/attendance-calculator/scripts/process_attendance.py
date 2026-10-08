#!/usr/bin/env python3
"""Generate the HR attendance summary from normalized attendance inputs."""

from __future__ import annotations

import argparse
import csv
import math
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple


OUTPUT_COLUMNS_TEMPLATE = [
    "部门",
    "职级",
    "姓名",
    "迟到分钟数（分钟）",
    "漏打卡次数",
    "迟到扣款情况",
    "漏打卡扣款情况",
    "考勤扣款情况（合计）",
    "本月事假时数（小时）",
    "累计事假时长（小时）",
    "累计加班结余（小时）",
    "病假（小时）",
    "育儿假",
    "奖励金",
    "本月应上班总小时数（h）",
    "入职时间",
    "离职时间",
    "上月病假时长",
    "病假天数合计(年/小时)",
    "{yy}年年假（天） 在职时间（{yy}.1.1-{yy}.12.31）",
    "上月年假剩余天数",
    "本月使用年假（天）",
    "{yy}年年假（天） 在职结余时间（{yy}.1.1-{yy}.12.31）",
    "考勤情况",
    "缺卡情况",
    "异常提示",
]


def read_table(path: Optional[str]) -> List[Dict[str, Any]]:
    if not path:
        return []
    table_path = Path(path)
    if not table_path.exists():
        raise FileNotFoundError(f"Input file not found: {table_path}")

    suffix = table_path.suffix.lower()
    if suffix in {".csv", ".txt"}:
        return read_csv(table_path)
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx(table_path)
    raise ValueError(f"Unsupported file type: {table_path.suffix}. Use CSV or XLSX.")


def read_csv(path: Path) -> List[Dict[str, Any]]:
    last_error: Optional[Exception] = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                return [normalize_row(row) for row in csv.DictReader(handle)]
        except UnicodeDecodeError as exc:
            last_error = exc
    raise UnicodeDecodeError("csv", b"", 0, 1, f"Could not decode {path}: {last_error}")


def read_xlsx(path: Path) -> List[Dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Reading XLSX requires openpyxl. Convert the file to CSV or install openpyxl.") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    records: List[Dict[str, Any]] = []
    for values in rows[1:]:
        row = {headers[index]: values[index] if index < len(values) else "" for index in range(len(headers)) if headers[index]}
        if any(str(value).strip() for value in row.values() if value is not None):
            records.append(normalize_row(row))
    return records


def normalize_row(row: Dict[str, Any]) -> Dict[str, Any]:
    return {str(key).strip(): value for key, value in row.items()}


def write_table(path: str, rows: List[Dict[str, Any]], columns: List[str]) -> None:
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    suffix = output_path.suffix.lower()
    if suffix == ".csv":
        with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=columns)
            writer.writeheader()
            writer.writerows(rows)
        return
    if suffix in {".xlsx", ".xlsm"}:
        write_xlsx(output_path, rows, columns)
        return
    raise ValueError("Output must be .csv or .xlsx.")


def write_xlsx(path: Path, rows: List[Dict[str, Any]], columns: List[str]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError as exc:
        raise RuntimeError("Writing XLSX requires openpyxl. Use a .csv output path or install openpyxl.") from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "综合汇总"
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(column, "") for column in columns])

    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, column in enumerate(columns, start=1):
        width = min(max(len(column) + 4, 12), 32)
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    workbook.save(path)


def parse_number(value: Any, default: float = 0.0) -> float:
    if value is None:
        return default
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return default
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return default
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group(0)) if match else default


def parse_int(value: Any, default: int = 0) -> int:
    return int(round(parse_number(value, float(default))))


def parse_date(value: Any) -> Optional[date]:
    if value is None or value == "":
        return None
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, (int, float)):
        try:
            base = datetime(1899, 12, 30)
            return (base + timedelta(days=float(value))).date()
        except Exception:
            return None
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def format_date(value: Optional[date]) -> str:
    return value.isoformat() if value else ""


def key_for(row: Dict[str, Any]) -> str:
    employee_id = str(row.get("employee_id", "") or "").strip()
    if employee_id:
        return f"id:{employee_id}"
    return f"name:{str(row.get('name', '') or '').strip()}"


def make_index(rows: Iterable[Dict[str, Any]]) -> Tuple[Dict[str, Dict[str, Any]], Dict[str, int]]:
    index: Dict[str, Dict[str, Any]] = {}
    counts: Dict[str, int] = {}
    for row in rows:
        key = key_for(row)
        if key in {"id:", "name:"}:
            continue
        counts[key] = counts.get(key, 0) + 1
        if key not in index:
            index[key] = row
    return index, counts


def add_years(day: date, years: int) -> date:
    try:
        return day.replace(year=day.year + years)
    except ValueError:
        return day.replace(month=2, day=28, year=day.year + years)


def ceil_to_half(value: float) -> float:
    if value <= 0:
        return 0.0
    return max(0.5, math.ceil(value * 2) / 2)


def annual_leave_entitlement(hire_date: Optional[date], year: int, month_start: Optional[date] = None) -> float:
    if not hire_date:
        return 0.0
    jan1 = date(year, 1, 1)
    one_year = add_years(hire_date, 1)
    three_years = add_years(hire_date, 3)
    five_years = add_years(hire_date, 5)

    if jan1 >= five_years:
        return 7.0
    if jan1 >= three_years:
        return 5.0
    if jan1 >= one_year:
        return 3.0
    if one_year.year == year and (month_start is None or month_start >= one_year):
        prorated = (12 - one_year.month) / 12 * 3
        return ceil_to_half(prorated)
    return 0.0


def split_birthdates(value: Any) -> List[date]:
    if value is None:
        return []
    parts = re.split(r"[,，;；/\n]+", str(value))
    dates = [parse_date(part.strip()) for part in parts if part.strip()]
    return [item for item in dates if item]


def child_year_window_for_month(birthdate: date, month_start: date) -> Optional[Tuple[date, date]]:
    third_birthday = add_years(birthdate, 3)
    if month_start > third_birthday:
        return None
    anchor_year = month_start.year
    anchor = birthdate.replace(year=anchor_year)
    if month_start < anchor:
        anchor = birthdate.replace(year=anchor_year - 1)
    end = add_years(anchor, 1)
    if anchor <= third_birthday and month_start <= third_birthday:
        return anchor, min(end, third_birthday)
    return None


def childcare_note(childcare_days: float, birthdates: List[date], month_start: date, prior_child_year_days: float) -> str:
    if childcare_days <= 0:
        return ""
    active_windows = [window for birthdate in birthdates if (window := child_year_window_for_month(birthdate, month_start))]
    notes: List[str] = [f"{clean_number(childcare_days)}天"]
    if not active_windows:
        notes.append("需复核：未找到3岁内子女")
    if childcare_days > 1:
        notes.append("需复核：本月超过1天")
    if prior_child_year_days + childcare_days > 10:
        notes.append("需复核：出生周年内超过10天")
    return "；".join(notes)


def clean_number(value: float) -> str:
    if abs(value - round(value)) < 0.00001:
        return str(int(round(value)))
    return f"{value:.2f}".rstrip("0").rstrip(".")


def detail_text(prefix: str, value: Any) -> str:
    text = str(value or "").strip()
    return f"{prefix}：{text}" if text else ""


def join_attendance_detail(late_detail: Any, personal_leave_detail: Any) -> str:
    parts = [detail_text("迟到", late_detail), detail_text("事假", personal_leave_detail)]
    return "\n".join(part for part in parts if part)


def output_columns(year: int) -> List[str]:
    yy = f"{year % 100:02d}"
    return [column.format(yy=yy) for column in OUTPUT_COLUMNS_TEMPLATE]


def build_summary(args: argparse.Namespace) -> Tuple[List[Dict[str, Any]], List[str]]:
    month = datetime.strptime(args.month, "%Y-%m").date()
    year = month.year
    columns = output_columns(year)

    attendance_rows = read_table(args.attendance)
    roster_index, roster_counts = make_index(read_table(args.roster))
    history_index, history_counts = make_index(read_table(args.history))
    overtime_index, overtime_counts = make_index(read_table(args.overtime))

    result_rows: List[Dict[str, Any]] = []
    for attendance in attendance_rows:
        key = key_for(attendance)
        roster = roster_index.get(key, {})
        history = history_index.get(key, {})
        overtime = overtime_index.get(key, {})
        exceptions: List[str] = []

        if roster_counts.get(key, 0) > 1 or history_counts.get(key, 0) > 1 or overtime_counts.get(key, 0) > 1:
            exceptions.append("员工标识重复，需复核")
        if not roster:
            exceptions.append("缺少花名册信息")

        attendance_type = str(attendance.get("attendance_type", "") or "").strip()
        needs_punch = attendance_type == "需要打卡"

        job_level = str(roster.get("job_level", "") or "").strip()
        rank = str(roster.get("rank", "") or job_level).strip()
        late_minutes = parse_int(attendance.get("late_minutes"))
        missed_punch_count = parse_int(attendance.get("missed_punch_count"))
        late_deduction: Any = ""
        if needs_punch:
            if job_level == "组员":
                late_deduction = late_minutes * -3
            elif job_level == "管理层":
                late_deduction = late_minutes * -5
            else:
                exceptions.append("缺少或无法识别职级，迟到扣款未计算")
        missed_deduction = max(0, missed_punch_count - 2) * -50 if needs_punch else 0
        total_deduction = (late_deduction if isinstance(late_deduction, int) else 0) + missed_deduction

        overtime_balance = parse_number(overtime.get("overtime_balance_hours"))
        overtime_to_personal_leave = abs(overtime_balance) if overtime_balance < 0 else 0.0

        monthly_personal_leave = parse_number(attendance.get("personal_leave_hours")) if needs_punch else 0.0
        monthly_personal_leave += overtime_to_personal_leave
        prior_personal_leave = parse_number(history.get("prior_personal_leave_hours"))
        cumulative_personal_leave = prior_personal_leave + monthly_personal_leave

        monthly_sick_leave = parse_number(attendance.get("sick_leave_hours")) if needs_punch else 0.0
        prior_sick_leave = parse_number(history.get("prior_sick_leave_hours_year"))
        ytd_sick_leave = prior_sick_leave + monthly_sick_leave

        hire_date = parse_date(roster.get("hire_date"))
        annual_entitlement = annual_leave_entitlement(hire_date, year, month)
        annual_used = parse_number(attendance.get("annual_leave_days"))
        prior_annual_remaining = parse_number(history.get("prior_annual_leave_remaining_days"), annual_entitlement)
        annual_remaining = max(0.0, prior_annual_remaining - annual_used)

        childcare_days = parse_number(attendance.get("childcare_leave_days"))
        birthdates = split_birthdates(roster.get("child_birthdates"))
        prior_childcare = parse_number(history.get("childcare_leave_days_in_child_year"))
        childcare = childcare_note(childcare_days, birthdates, month, prior_childcare)

        department = str(attendance.get("department", "") or roster.get("department", "") or "").strip()
        missed_detail = detail_text("缺卡", attendance.get("missed_punch_detail"))
        attendance_detail = join_attendance_detail(attendance.get("late_detail"), attendance.get("personal_leave_detail"))
        if overtime_to_personal_leave:
            addition = f"加班结余为负转事假{clean_number(overtime_to_personal_leave)}小时"
            attendance_detail = "\n".join(part for part in [attendance_detail, f"事假：{addition}"] if part)

        row = {
            "部门": department,
            "职级": rank,
            "姓名": str(attendance.get("name", "") or roster.get("name", "") or "").strip(),
            "迟到分钟数（分钟）": late_minutes if needs_punch else "",
            "漏打卡次数": missed_punch_count if needs_punch else "",
            "迟到扣款情况": late_deduction,
            "漏打卡扣款情况": missed_deduction if needs_punch else "",
            "考勤扣款情况（合计）": total_deduction if needs_punch else "",
            "本月事假时数（小时）": clean_number(monthly_personal_leave) if needs_punch or overtime_to_personal_leave else "",
            "累计事假时长（小时）": clean_number(cumulative_personal_leave) if needs_punch or overtime_to_personal_leave else "",
            "累计加班结余（小时）": clean_number(overtime_balance),
            "病假（小时）": clean_number(monthly_sick_leave) if needs_punch else "",
            "育儿假": childcare,
            "奖励金": "",
            "本月应上班总小时数（h）": clean_number(parse_number(attendance.get("required_work_hours"), args.paid_days * args.hours_per_day)),
            "入职时间": format_date(hire_date),
            "离职时间": "",
            "上月病假时长": clean_number(prior_sick_leave),
            "病假天数合计(年/小时)": clean_number(ytd_sick_leave),
            columns[19]: clean_number(annual_entitlement),
            "上月年假剩余天数": clean_number(prior_annual_remaining),
            "本月使用年假（天）": clean_number(annual_used),
            columns[22]: clean_number(annual_remaining),
            "考勤情况": attendance_detail,
            "缺卡情况": missed_detail,
            "异常提示": "；".join(exceptions),
        }
        result_rows.append(row)
    return result_rows, columns


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate HR attendance summary from normalized inputs.")
    parser.add_argument("--month", required=True, help="Calculation month in YYYY-MM format.")
    parser.add_argument("--paid-days", required=True, type=float, help="Actual paid workdays in the month.")
    parser.add_argument("--hours-per-day", default=8.0, type=float, help="Work hours per day. Default: 8.")
    parser.add_argument("--attendance", required=True, help="Normalized monthly attendance CSV/XLSX.")
    parser.add_argument("--roster", required=True, help="Normalized roster CSV/XLSX.")
    parser.add_argument("--history", help="Normalized history CSV/XLSX.")
    parser.add_argument("--overtime", help="Normalized overtime CSV/XLSX.")
    parser.add_argument("--output", required=True, help="Output CSV/XLSX path.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    rows, columns = build_summary(args)
    write_table(args.output, rows, columns)
    print(f"Wrote {len(rows)} rows to {args.output}")


if __name__ == "__main__":
    main()
